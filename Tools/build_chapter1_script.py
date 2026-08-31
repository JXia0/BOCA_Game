from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\GameProjects\BOCA_Game\Docs\BOCA_Chapter1_Oil_on_the_Tide_CN.docx")

BLUE = "1F4D78"
MID_BLUE = "2E74B5"
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF5D6"
PALE_RED = "FCE8E6"
PALE_GREEN = "E8F3EA"
GRAY = "666666"
LIGHT_GRAY = "F2F4F7"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def apply_font(run, latin="Calibri", cjk="Microsoft YaHei", size=11, bold=None, color=BLACK, italic=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cjk)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MID_BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run_properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    run_properties.append(size)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(run_properties)
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_para(doc, text="", style=None, before=0, after=6, line=1.25, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    if text:
        r = p.add_run(text)
        apply_font(r)
    return p


def add_label_para(doc, label, text, fill=None):
    p = doc.add_paragraph()
    if fill:
        set_paragraph_shading(p, fill)
        p.paragraph_format.left_indent = Inches(0.08)
        p.paragraph_format.right_indent = Inches(0.08)
    r = p.add_run(label)
    apply_font(r, bold=True, color=BLUE)
    r = p.add_run(text)
    apply_font(r)
    p.paragraph_format.space_before = Pt(2 if fill else 0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_dialogue(doc, speaker, line, direction=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(f"{speaker}：")
    apply_font(r, bold=True, color=BLUE)
    r = p.add_run(line)
    apply_font(r)
    if direction:
        r = p.add_run(f" 〔{direction}〕")
        apply_font(r, size=9.5, italic=True, color=GRAY)
    return p


def add_direction(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"【画面/演出】{text}")
    apply_font(r, size=10, italic=True, color=GRAY)
    return p


def add_ui(doc, text, fill=PALE_BLUE):
    return add_label_para(doc, "【UI】", text, fill=fill)


def add_action(doc, text, result=None):
    p = doc.add_paragraph(style="Action Step")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    apply_font(r, bold=True, color=BLUE)
    if result:
        r = p.add_run(f" → {result}")
        apply_font(r)
    return p


def add_choice_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["玩家选择", "即时反馈", "流程结果"]
    for index, text in enumerate(headers):
        cell = table.cell(0, index)
        set_cell_shading(cell, BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        apply_font(r, size=10, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for choice, feedback, outcome, color in rows:
        cells = table.add_row().cells
        for idx, text in enumerate((choice, feedback, outcome)):
            set_cell_shading(cells[idx], color)
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            apply_font(r, size=9.5, bold=(idx == 0))
    set_table_geometry(table, [2160, 3600, 3600])
    add_para(doc, "", after=2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        apply_font(
            run,
            size={1: 16, 2: 13, 3: 12}[level],
            bold=True,
            color=MID_BLUE if level < 3 else BLUE,
        )
    return p


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, before, after in ((1, 16, 18, 10), (2, 13, 14, 7), (3, 12, 10, 5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(MID_BLUE if level < 3 else BLUE)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Action Step" not in styles:
        action_style = styles.add_style("Action Step", WD_STYLE_TYPE.PARAGRAPH)
    else:
        action_style = styles["Action Step"]
    action_style.base_style = normal
    action_style.paragraph_format.left_indent = Inches(0.375)
    action_style.paragraph_format.first_line_indent = Inches(-0.188)
    action_style.paragraph_format.space_after = Pt(4)
    action_style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("BOCA｜章节制作脚本")
    apply_font(r, size=9, bold=True, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("内部制作稿｜事实、玩法与安全审核后方可锁稿")
    apply_font(r, size=8.5, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("BOCA 第一章")
    apply_font(r, size=11, bold=True, color=MID_BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("油潮之下")
    apply_font(r, size=28, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Oil on the Tide｜企鹅救助与“耐心的关怀”")
    apply_font(r, size=14, color=GRAY)

    meta = doc.add_table(rows=1, cols=2)
    meta.style = "Table Grid"
    for index, text in enumerate(("字段", "章节规格")):
        set_cell_shading(meta.cell(0, index), BLUE)
        r = meta.cell(0, index).paragraphs[0].add_run(text)
        apply_font(r, size=10, bold=True, color=WHITE)
    set_repeat_table_header(meta.rows[0])
    metadata = [
        ("状态", "第一版制作脚本｜等待野生动物救助与历史事实审核"),
        ("目标时长", "10–12分钟；无计时模式下不限时"),
        ("玩家身份", "虚构的当地青年志愿协助者“你”"),
        ("核心课程", "关怀不是立刻动手，而是观察、求助、等待专业判断"),
        ("章节奖励", "CARE｜关怀徽章碎片"),
        ("主要玩法", "现场观察、秩序排列、抽象护理协助、恢复记录判断"),
    ]
    for label, value in metadata:
        cells = meta.add_row().cells
        set_cell_shading(cells[0], PALE_BLUE)
        r = cells[0].paragraphs[0].add_run(label)
        apply_font(r, size=10, bold=True, color=BLUE)
        r = cells[1].paragraphs[0].add_run(value)
        apply_font(r, size=10)
    set_table_geometry(meta, [1700, 7660])

    add_para(doc, "")
    add_label_para(
        doc,
        "安全底线｜",
        "普通玩家不得自行触碰、喂食、清洗或放归油污野生动物。所有近距离处置均由专业机构执行；玩家只在明确指令下进行低风险协助。",
        fill=PALE_RED,
    )

    add_heading(doc, "1. 事实与虚构边界", 1)
    add_heading(doc, "已核实的历史锚点", 2)
    for text in (
        "2011年，巴西海岸发现一只满身油污、虚弱饥饿的麦哲伦企鹅。",
        "这只企鹅后来被称为Dindim，在康复后被放归海洋。",
        "后续报道描述Dindim再次返回救助者所在岛屿；具体迁徙距离、动机和长期行为不作确定性解释。",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        apply_font(r)

    add_heading(doc, "为游戏虚构或压缩的内容", 2)
    for text in (
        "玩家角色、专家玛丽娜、游客和工作人员均为虚构。",
        "场景地点、对白、护理中心布局、日志数据和时间线为教学目的而压缩。",
        "清洗互动是抽象化的专业流程，不包含化学品、浓度、温度或可供现实模仿的操作配方。",
        "真实人物若昂不作为可操作角色；其姓名只在历史回顾中出现。",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        apply_font(r)

    add_heading(doc, "2. 章节节奏", 1)
    timing = doc.add_table(rows=1, cols=4)
    timing.style = "Table Grid"
    for idx, text in enumerate(("段落", "预计时长", "玩家任务", "学习目标")):
        set_cell_shading(timing.cell(0, idx), BLUE)
        r = timing.cell(0, idx).paragraphs[0].add_run(text)
        apply_font(r, size=9.5, bold=True, color=WHITE)
    timing_rows = [
        ("档案馆引入", "0:45", "进入记忆", "关怀碎片主题"),
        ("海滩发现", "2:15", "观察并建立安静区域", "先观察、再求助"),
        ("专业转运", "1:45", "排列安全步骤", "普通人与专业人员的边界"),
        ("护理中心", "2:30", "抽象护理协助", "稳定优先，不追求立刻洗净"),
        ("恢复评估", "2:00", "比较日志并申请评估", "外表干净不等于可以放归"),
        ("放归与回顾", "1:30", "反思并恢复碎片", "承认不确定性"),
    ]
    for row in timing_rows:
        cells = timing.add_row().cells
        for idx, text in enumerate(row):
            r = cells[idx].paragraphs[0].add_run(text)
            apply_font(r, size=9.25, bold=(idx == 0))
    set_repeat_table_header(timing.rows[0])
    set_table_geometry(timing, [2200, 1200, 3000, 2960])

    doc.add_page_break()
    add_heading(doc, "3. 正式章节脚本", 1)

    add_heading(doc, "场景0｜档案馆：缺失的关怀", 2)
    add_label_para(doc, "时间｜", "约45秒　　地点｜深海档案馆　　事实标签｜完全虚构", fill=LIGHT_GRAY)
    add_direction(doc, "深蓝色档案馆中，一枚缺失四角的奖章悬在玻璃台上。第一处空位泛起微弱的琥珀色光。远处传来海浪与短促的鸟鸣。")
    add_dialogue(doc, "档案管理员", "第一块碎片不是奖励勇敢。它记录的是另一种更难的事——在想立刻帮忙的时候，先停下来，确认什么才真正安全。")
    add_dialogue(doc, "档案管理员", "这段记忆来自2011年的巴西。真实事件会保留；为了让你参与其中，角色和时间线经过了压缩。")
    add_ui(doc, "记忆主题：CARE｜关怀　　进入记忆")
    add_action(doc, "玩家按下“进入记忆”", "画面被一层海水般的蓝光覆盖，转入海滩。")

    add_heading(doc, "场景1｜海滩：先不要碰", 2)
    add_label_para(doc, "时间｜", "约2分15秒　　地点｜巴西近海岛屿沙滩　　事实标签｜地点与人物虚构，动物状态来自案例背景", fill=LIGHT_GRAY)
    add_direction(doc, "清晨。潮线附近散落海草和少量油污。一个黑白色身影伏在岩石旁，羽毛被油黏成不规则的深色块。它呼吸急促，但仍会轻微转头。")
    add_dialogue(doc, "游客", "它是不是冷？我把外套盖上去，再喂它一点鱼？")
    add_ui(doc, "新任务：观察现场。使用“专注”找出4项重要信息。")
    for action, result in (
        ("观察企鹅", "记录：虚弱、羽毛油污、仍有呼吸反应。"),
        ("观察潮线", "记录：海水正在上涨，但动物当前无法正常游泳。"),
        ("观察人群与宠物", "记录：游客靠近；一只未牵绳的狗正在吠叫。"),
        ("观察地面", "记录：附近有油污和尖锐垃圾，接近可能伤害人和动物。"),
    ):
        add_action(doc, action, result)

    add_ui(doc, "选择你的第一步。", fill=PALE_GOLD)
    add_choice_table(
        doc,
        [
            ("直接抱起企鹅", "专业机构尚未到场；接触可能加重应激，也可能伤到你。", "解释后回到选择。", PALE_RED),
            ("把企鹅推回海里", "虚弱动物可能再次搁浅或溺水。", "解释后回到选择。", PALE_RED),
            ("保持距离并联系救援机构", "你提供了位置、动物状态和现场风险。", "正确路径；任务继续。", PALE_GREEN),
        ],
    )
    add_dialogue(doc, "你", "这里有一只疑似受油污影响的企鹅。它还在呼吸，位置在北侧岩礁旁。附近有人群、狗和油污，我会保持距离，等你们指示。")
    add_dialogue(doc, "救援热线", "收到。请不要触碰、喂食、清洗或把它送回海里。先让周围安静下来，并记住它的位置。救援人员正在赶来。")

    add_ui(doc, "小游戏：安静区域。拖动3个提示标志到正确位置。")
    for action, result in (
        ("把“请勿靠近”放在人群前方", "游客退到安全距离。"),
        ("把“宠物请牵绳”放在入口", "狗被带离动物附近。"),
        ("把“油污区域”标在污染边缘", "避免他人踩入并扩散污染。"),
    ):
        add_action(doc, action, result)
    add_dialogue(doc, "游客", "我们真的什么都不做吗？")
    add_dialogue(doc, "你", "我们正在做：不让情况变得更糟，并让专业人员能安全接手。")

    add_heading(doc, "场景2｜海滩：专业人员接手", 2)
    add_label_para(doc, "时间｜", "约1分45秒　　事实标签｜人物和具体步骤虚构；安全原则来自官方指导", fill=LIGHT_GRAY)
    add_direction(doc, "一辆带有野生动物救援标识的车辆停下。专家玛丽娜和两名工作人员穿戴防护装备，携带遮挡板、转运箱和记录工具。")
    add_dialogue(doc, "玛丽娜", "谢谢你没有移动它。先告诉我：什么时候发现的？有没有人碰过、喂过或清洗过？")
    add_dialogue(doc, "你", "大约十分钟前发现。没有人接触。我们把狗和游客带远，也标出了油污区域。")
    add_dialogue(doc, "玛丽娜", "很好。现在由我们检查和转运。你可以帮我核对记录，但不要靠近动物。")
    add_ui(doc, "顺序谜题：把5张行动卡排列成安全流程。")
    add_action(doc, "正确顺序", "观察记录 → 联系专业机构 → 隔离风险 → 专业检查与转运 → 护理后再评估放归")
    add_ui(doc, "错误卡示例：立即喂食、用清洁剂冲洗、直接放回海里。放入流程时会解释原因并弹回卡槽。", fill=PALE_GOLD)
    add_dialogue(doc, "玛丽娜", "油污不只是看起来脏。它会破坏羽毛的防水和保温能力。清洗前，动物是否稳定同样重要。")
    add_direction(doc, "工作人员用遮挡板降低视觉刺激，在专业操作下将企鹅转入通风转运箱。镜头保持中景，不展示抓取细节。")

    add_heading(doc, "场景3｜护理中心：干净不是第一步", 2)
    add_label_para(doc, "时间｜", "约2分30秒　　地点｜虚构的野生动物护理中心　　事实标签｜流程抽象化，不可作为现实教程", fill=LIGHT_GRAY)
    add_direction(doc, "室内光线柔和。企鹅位于专业护理区，玩家隔着观察窗查看三项状态：应激、体力、羽毛污染。所有直接操作由工作人员完成。")
    add_dialogue(doc, "玛丽娜", "很多人看到油污，第一反应是立刻洗。但如果动物还不稳定，清洗本身也可能成为负担。")
    add_ui(doc, "小游戏：护理协助。根据状态面板选择下一阶段；不含配方与现实操作细节。")
    add_choice_table(
        doc,
        [
            ("马上进入清洗阶段", "体力仍低，应激较高。", "解释后回到状态面板。", PALE_RED),
            ("先等待专业人员完成稳定评估", "体力与应激状态逐渐改善。", "解锁下一阶段。", PALE_GREEN),
        ],
    )
    add_direction(doc, "时间压缩蒙太奇：工作人员记录、检查并准备护理；不展示药物、喂食或清洁剂细节。")
    add_ui(doc, "抽象清洗阶段：玩家只负责观察3个指标并在工作人员提示时按下“暂停”。")
    for action, result in (
        ("污染残留下降，但应激上升", "选择“暂停”，让工作人员重新评估。"),
        ("应激稳定，专业人员确认可继续", "选择“继续下一小段”。"),
        ("羽毛表面已无明显油污", "不能直接判定可放归；进入恢复观察。"),
    ):
        add_action(doc, action, result)
    add_dialogue(doc, "你", "看起来已经干净了。这样就结束了吗？")
    add_dialogue(doc, "玛丽娜", "还没有。羽毛能否重新防水、它能否正常游泳和维持状态，比‘看起来干净’更重要。")

    add_heading(doc, "场景4｜恢复记录：什么时候才算准备好", 2)
    add_label_para(doc, "时间｜", "约2分钟　　事实标签｜日志数据为虚构教学数据", fill=LIGHT_GRAY)
    add_ui(doc, "打开恢复日志。比较三个检查日，不设置隐藏分数。")
    logs = doc.add_table(rows=1, cols=5)
    logs.style = "Table Grid"
    for idx, text in enumerate(("记录", "精神/体力", "羽毛防水", "游泳表现", "专业结论")):
        set_cell_shading(logs.cell(0, idx), BLUE)
        r = logs.cell(0, idx).paragraphs[0].add_run(text)
        apply_font(r, size=9, bold=True, color=WHITE)
    for row in (
        ("A", "仍虚弱", "未通过", "未测试", "继续护理"),
        ("B", "改善", "部分恢复", "短时稳定", "继续观察"),
        ("C", "稳定", "通过专业检查", "持续稳定", "可申请放归评估"),
    ):
        cells = logs.add_row().cells
        for idx, text in enumerate(row):
            r = cells[idx].paragraphs[0].add_run(text)
            apply_font(r, size=9, bold=(idx == 0))
    set_repeat_table_header(logs.rows[0])
    set_table_geometry(logs, [900, 1800, 2000, 1900, 2760])
    add_ui(doc, "选择：根据哪条记录申请专业放归评估？", fill=PALE_GOLD)
    add_choice_table(
        doc,
        [
            ("记录A：油污少了", "单看外表不足；体力和功能尚未恢复。", "回到日志。", PALE_RED),
            ("记录B：已经能短时游泳", "仍需观察稳定性与羽毛功能。", "回到日志。", PALE_RED),
            ("记录C：多项功能稳定", "你提交的是“申请评估”，最终决定仍由专业人员作出。", "进入放归评估。", PALE_GREEN),
        ],
    )
    add_dialogue(doc, "你", "我选择记录C，并申请评估——不是宣布它一定可以放归。")
    add_dialogue(doc, "玛丽娜", "正确。真正负责的决定，也包括承认自己还不能确定。")

    add_heading(doc, "场景5｜海边：放归不是保证", 2)
    add_label_para(doc, "时间｜", "约1分30秒　　事实标签｜Dindim康复、放归和被报道返回为历史锚点；现场细节虚构", fill=LIGHT_GRAY)
    add_direction(doc, "清晨海面平稳。专业团队完成最后评估。玩家站在标线外。转运箱朝向海面打开，镜头保持远景。企鹅停顿片刻后走向水边，进入浪中。")
    add_dialogue(doc, "你", "我们成功了吗？")
    add_dialogue(doc, "玛丽娜", "我们完成了今天能完成的部分。放归很重要，但它不会让未来变得确定。")
    add_direction(doc, "企鹅身影变小，最终消失在水面。画面不拟人化，不让企鹅回头挥别。")
    add_ui(doc, "历史回顾：2011年，一只后来被称为Dindim的麦哲伦企鹅在巴西被发现时满身油污、十分虚弱。它康复后重返海洋。后续报道描述它再次返回救助者所在的岛屿。关于迁徙距离和返回动机，游戏不作确定解释。")
    add_ui(doc, "现实安全提示：如果发现受油污影响的野生动物，请保持距离、记录位置并联系当地野生动物或环境机构。不要自行触碰、喂食、清洗或放归。", fill=PALE_RED)

    add_heading(doc, "场景6｜档案馆：恢复CARE碎片", 2)
    add_direction(doc, "画面回到深海档案馆。琥珀色碎片中浮现企鹅与一圈干净水纹的轮廓。")
    add_dialogue(doc, "档案管理员", "关怀不是把生命握在手里。它也可以是退后一步、听从专业判断，并耐心等到真正适合的时刻。")
    add_ui(doc, "CARE碎片已恢复。章节完成不要求完美选择；所有错误选择都已通过解释和回退修正。", fill=PALE_GREEN)
    add_choice_table(
        doc,
        [
            ("关怀意味着马上行动", "行动必须建立在安全和专业判断上。", "记录反思，不影响奖励。", PALE_GOLD),
            ("关怀意味着耐心地做正确的事", "你将观察、求助、等待和评估连接起来。", "获得“耐心的关怀”掌握标记。", PALE_GREEN),
        ],
    )

    doc.add_page_break()
    add_heading(doc, "4. 制作与验收说明", 1)
    add_heading(doc, "程序实现要求", 2)
    for text in (
        "所有选择使用失败回退，不设置永久失败或错过CARE碎片的状态。",
        "所有对白、UI、安全提示、事实卡和来源说明均外置为本地化数据，不得硬编码。",
        "“护理协助”仅操作抽象状态，不出现真实清洗配方、药物、喂食量或可模仿步骤。",
        "保存点：进入海滩、专业人员到场、护理中心、恢复日志、放归前。",
        "辅助模式：高亮所有观察点；自动排列流程卡；状态指标只增不减；关闭时间压力。",
    ):
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(text)
        apply_font(r)

    add_heading(doc, "美术与音频要求", 2)
    for text in (
        "企鹅痛苦通过姿态、呼吸、羽毛颜色和环境呈现；禁止血腥、解剖特写或夸张哭泣。",
        "油污不能被画成可爱花纹；必须清楚表现为影响羽毛结构和行动的污染。",
        "海滩、建筑、服装和救援装备需要巴西地区参考审核，不能画成南极基地。",
        "企鹅不向玩家鞠躬、挥手、流泪或回头告别。",
        "所有重要音效同时提供字幕或视觉提示；鸟鸣不作为唯一线索。",
    ):
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(text)
        apply_font(r)

    add_heading(doc, "剧本验收条件", 2)
    checks = (
        "首次玩家可在无外部指导下完成章节。",
        "错误选择解释具体风险并返回当前决策，不重开章节。",
        "至少一次明确区分“看起来干净”与“功能恢复”。",
        "至少两次明确告诉玩家不要自行触碰、喂食、清洗或放归。",
        "结尾不宣称Dindim必然出于感恩返回，也不确定其完整迁徙路线。",
        "所有历史信息附来源；所有虚构角色和压缩时间线有标记。",
    )
    for text in checks:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(text)
        apply_font(r)

    add_heading(doc, "5. 来源说明", 1)
    sources = [
        ("S1｜中国新闻网，2016-03-10", "Dindim被报道于2011年在巴西获救、康复放归，并在之后返回岛屿。", "https://www.chinanews.com.cn/m/gj/2016/03-10/7791193.shtml"),
        ("S2｜巴西环境与可再生自然资源研究所（IBAMA）", "发现油污动物时应立即联系机构；不要触碰、放回海里、使用清洁产品或喂食。", "https://www.gov.br/ibama/pt-br/acesso-a-informacao/perguntas-frequentes/manchas-de-oleo"),
        ("S3｜Rio das Ostras市政府", "麦哲伦企鹅会迁徙至巴西海岸；发现时应保持距离、隔离人和宠物并联系救援机构。", "https://www.riodasostras.rj.gov.br/pinguim-e-resgatado-com-vida/"),
        ("S4｜Marine Pollution Bulletin / PubMed", "专业康复与放归是减轻油污影响的一种保护策略，但长期结果需要持续监测。", "https://pubmed.ncbi.nlm.nih.gov/22683106/"),
    ]
    for label, description, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label)
        apply_font(r, size=10, bold=True, color=BLUE)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(description + " ")
        apply_font(r, size=10)
        add_hyperlink(p, "打开来源", url)
        p.paragraph_format.space_after = Pt(7)

    add_label_para(
        doc,
        "发布前必做｜",
        "由巴西当地野生动物救助专业人员复核救援表述；由历史/编辑人员复核Dindim报道；由教育与无障碍人员进行10岁以上玩家测试。",
        fill=PALE_GOLD,
    )

    core = doc.core_properties
    core.title = "BOCA 第一章：油潮之下"
    core.subject = "Dindim企鹅救助章节制作脚本"
    core.author = "BOCA Project Team"
    core.keywords = "BOCA, Dindim, penguin, wildlife rescue, game script"
    core.comments = "Source-grounded production draft. Fiction and safety boundaries are labeled."

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
