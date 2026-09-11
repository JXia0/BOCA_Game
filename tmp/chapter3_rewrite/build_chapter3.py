from pathlib import Path
import os
import stat
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\GameProjects\BOCA_Game")
OUTPUT = ROOT / "output" / "docx" / "第三章_暗礁忧鸣_统一标准版.docx"

PURPLE = RGBColor(0xA0, 0x2B, 0x93)
RED = RGBColor(0xEE, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x9B, 0x45)
GRAY = RGBColor(0x66, 0x66, 0x66)
BLACK = RGBColor(0x00, 0x00, 0x00)


def apply_font(run, size=10.5):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    fonts.set(qn("w:cs"), "Aptos")


def format_run(run, fmt):
    apply_font(run, fmt.get("size", 10.5))
    run.bold = bool(fmt.get("bold"))
    run.italic = bool(fmt.get("italic"))
    if fmt.get("color") is not None:
        run.font.color.rgb = fmt["color"]
    if fmt.get("highlight") == "yellow":
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    elif fmt.get("highlight") == "turquoise":
        run.font.highlight_color = WD_COLOR_INDEX.TURQUOISE


def paragraph(doc, parts, *, style=None, before=0, after=5, keep=False, page_break=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.keep_together = keep
    p.paragraph_format.page_break_before = page_break
    for text, fmt in parts:
        run = p.add_run(text)
        format_run(run, fmt)
    return p


def plain(doc, text, **kwargs):
    return paragraph(doc, [(text, {})], **kwargs)


def heading(doc, text, level=1, page_break=False):
    p = paragraph(
        doc,
        [(text, {"bold": True, "size": 12.5 if level == 1 else 11})],
        style="Heading 1" if level == 1 else "Heading 2",
        before=10 if level == 1 else 6,
        after=5,
        keep=True,
        page_break=page_break,
    )
    p.paragraph_format.keep_with_next = True
    return p


def action(doc, text):
    return paragraph(doc, [(text, {"highlight": "yellow"})], after=4)


def choice(doc, label, text):
    return paragraph(
        doc,
        [(label + "  ", {"bold": True, "highlight": "turquoise"}), (text, {})],
        after=4,
    )


def colored(doc, label, text, color=PURPLE):
    return paragraph(
        doc,
        [(label + "  ", {"bold": True, "color": color}), (text, {"color": color})],
        after=5,
    )


def production(doc, text):
    return paragraph(
        doc,
        [
            ("[制作说明]  ", {"bold": True, "color": GRAY}),
            (text, {"italic": True, "color": GRAY}),
        ],
        after=5,
    )


def speaker(doc, name, text):
    return paragraph(doc, [(name + "：", {"bold": True}), (text, {})], after=4)


def hyperlink(paragraph_obj, text, url):
    rid = paragraph_obj.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "A02B93")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    link.append(run)
    paragraph_obj._p.append(link)


def source(doc, title, url, use):
    p = paragraph(doc, [(title + "  ", {"bold": True})], after=1)
    hyperlink(p, url, url)
    p.paragraph_format.keep_with_next = True
    paragraph(doc, [(use, {"color": GRAY, "size": 9.5})], after=5, keep=True)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    apply_font(run, 9)
    run.font.color.rgb = GRAY
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        ppr = style._element.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)
    doc.styles["Title"].font.size = Pt(18)
    doc.styles["Title"].font.bold = True

    doc.core_properties.title = "BOCA 第三章 暗礁忧鸣 统一标准版"
    doc.core_properties.subject = "围头湾中华白海豚栖息地调查互动游戏剧本"
    doc.core_properties.author = "BOCA Project Team"

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("第三章 暗礁忧鸣")
    apply_font(run, 18)
    run.bold = True
    paragraph(
        doc,
        [("围头湾中华白海豚栖息地调查互动剧本 统一标准版", {"italic": True, "color": GRAY, "size": 10})],
        after=7,
    )
    paragraph(
        doc,
        [
            ("格式说明  ", {"bold": True}),
            ("黄色高亮", {"highlight": "yellow"}),
            ("表示玩家操作  ", {}),
            ("青色高亮", {"highlight": "turquoise"}),
            ("表示玩家选择  ", {}),
            ("绿色", {"color": GREEN}),
            ("表示证据允许的结论  ", {}),
            ("紫色", {"color": PURPLE}),
            ("表示系统或研究信息  ", {}),
            ("红色", {"color": RED}),
            ("表示越界推断或风险提示", {}),
        ],
        after=7,
    )
    production(
        doc,
        "本章以2025年围头湾泊位工程炸礁争议为现实背景。玩家、潮汐观察站、渔民对话、现场观察和证据界面均为游戏化创作；真实机构、工程与公开事件只在来源能够支持的范围内呈现。地图与目击点首次出现时必须显示“游戏化模拟数据”，不得伪装成真实监测成果。",
    )
    production(
        doc,
        "现实信息更新截至2026年9月9日。诉讼状态、工程状态、机构名称和对外报告渠道须放在可更新的数据层，发行前再次核对。",
    )

    heading(doc, "角色与任务")
    plain(
        doc,
        "玩家是民间环境观察组织“潮汐观察站”的田野调查志愿者。组织不隶属于政府，也没有执法权。玩家负责保存公开资料、在公共区域观察、经同意记录居民口述，并向组织项目负责人提交内部证据包。",
    )
    colored(
        doc,
        "[权限]",
        "玩家可以记录亲眼所见、保存公开网页、提出核实请求，也可以标注“不知道”。玩家不能进入施工控制区、调取非公开文件、宣布工程违法、认定动物伤亡或命令停工。",
    )
    colored(
        doc,
        "[核心目标]",
        "把事实、证人陈述、机构主张、一般风险知识和未知事项分开，让每条结论都能回到它真正依赖的来源。",
    )

    heading(doc, "场景一 施工公告", page_break=True)
    plain(doc, "时间卡：2025年5月12日 19:30")
    plain(doc, "地点：潮汐观察站办公室。窗外已经天黑，桌面上放着相机、录音笔、反光背心和空白观察表。")
    speaker(
        doc,
        "项目负责人",
        "志愿者刚发来一张海事公告截图。公告涉及围头湾一项泊位工程，预告明天进行水下爆破作业。先核对原网页，再决定现场怎么记录。",
    )
    action(doc, "打开公告网页，保存完整页面、网址、发布时间和截图时间。")
    colored(
        doc,
        "[记录 E01]",
        "公开公告可以证明有关单位曾预告作业，但不能单独证明作业后来实际发生，也不能证明白海豚受到影响。",
        GREEN,
    )
    action(doc, "在地图上标出公共岸线观察点、施工控制区和撤离路线。")
    colored(
        doc,
        "[安全检查]",
        "观察点必须位于公开区域。相机使用长焦镜头；玩家不登施工船、不翻越围栏，也不使用无人机进入受限空域。",
    )
    choice(doc, "选择 A", "把任务标题写成“爆破将伤害白海豚”。")
    colored(
        doc,
        "[系统结果]",
        "“将伤害”把风险写成确定结果。调查可信度减1；这句话会留在草稿中，等待最终审核。",
        RED,
    )
    choice(doc, "选择 B", "把任务标题写成“核实爆破计划、现场情况与白海豚活动线索”。")
    colored(doc, "[系统结果]", "标题没有预设调查结论。调查可信度加1。", GREEN)
    action(doc, "确认装备清单，结束当天准备。")

    heading(doc, "场景二 公共岸线")
    plain(doc, "时间卡：2025年5月13日 18:35")
    plain(doc, "海面仍有施工船，岸上能听见发动机和金属碰撞声。玩家在公共岸线停下，没有进入工程区域。")
    action(doc, "拍摄一张包含海面、施工船和固定地标的全景照片。")
    colored(
        doc,
        "[记录 E02]",
        "时间：18:38。地点：围头湾石井作业区附近公共岸线。观察：远处可见施工船。局限：照片不能显示水下情况，也不能单独证明正在爆破。",
        GREEN,
    )
    plain(doc, "一名当地渔民收起工具，准备离开。")
    speaker(
        doc,
        "玩家",
        "您好，我是潮汐观察站的志愿者。可以问您几个问题并做文字记录吗？如果记不清，请直接说不确定。您的姓名可以不写。",
    )
    speaker(doc, "渔民", "可以。我只能讲自己听到和看到的。")
    colored(
        doc,
        "[采访状态]",
        "对方已知情并同意。记录匿名；不录制面部；方言版本若加入游戏，须由本地说话者审核。",
    )
    speaker(doc, "玩家", "您今天注意到什么？")
    speaker(
        doc,
        "渔民",
        "下午听见过两声很闷的响声。我人在岸上，看不到海面下面，所以不能说那一定是什么。",
    )
    speaker(doc, "玩家", "您今天看到白海豚了吗？")
    speaker(
        doc,
        "渔民",
        "今天没看清。我过去在这片海见过白色的背鳍，但没记日期，也不知道它们有没有受伤。",
    )
    action(doc, "把“亲耳听到”“没有看见爆破”“过去曾见疑似白海豚”“日期不详”拖入四个独立字段。")
    colored(
        doc,
        "[记录 E03]",
        "这是单一证人的口述。它能保存个人经历，不能独立确认声源、爆破次数、动物种类或因果关系。",
        GREEN,
    )
    choice(doc, "选择 A", "记录为“渔民证实爆破后白海豚受伤”。")
    colored(
        doc,
        "[系统结果]",
        "证人没有看到爆破，也没有报告动物受伤。该表述篡改了证词。调查可信度减2；最终审核将要求改正。",
        RED,
    )
    choice(doc, "选择 B", "照原意记录，并把声源、日期与物种识别列为待核实。")
    colored(doc, "[系统结果]", "观察与推测已分开。调查可信度加1。", GREEN)
    plain(doc, "太阳接近海平线。远处短暂出现几道浅色背鳍，随后消失。玩家没有靠近。")
    action(doc, "用长焦镜头记录短视频，同时填写方向、距离估计、天气和拍摄时间。")
    colored(
        doc,
        "[记录 E04]",
        "画面可能记录到中华白海豚，但需要有资质人员复核。一次目击不能代表完整分布，也不能证明动物因当天施工受伤。",
        GREEN,
    )

    heading(doc, "场景三 核对公开信息", page_break=True)
    plain(doc, "时间卡：2025年5月16日 20:10")
    plain(doc, "玩家回到办公室。新闻报道、机构回应和项目公开文件同时出现在资料库中。")
    action(doc, "打开5月15日至16日的两篇报道，逐项核对报道者、受访机构、原话范围和发布时间。")
    colored(
        doc,
        "[资料 E05]",
        "澎湃新闻5月16日报道称，福建省海洋局渔业处工作人员表示：5月13日进行了两次炸礁作业，当晚被叫停；5月14日有关人员到场评估，并要求重新遴选施工工艺、完善中华白海豚保护专题论证。",
    )
    colored(
        doc,
        "[证据边界]",
        "这是媒体对部门工作人员说法的转述，不应改写成玩家亲见、法院认定或一份未取得的正式通报。",
        RED,
    )
    colored(
        doc,
        "[资料 E06]",
        "红星新闻5月15日报道称，当时暂未收到中华白海豚伤亡报告。“未收到报告”不等于已经完成全面监测，也不等于可以证明没有影响。",
    )
    colored(
        doc,
        "[时间线]",
        "5月13日作业、当晚暂停、5月14日现场评估，均发生在玩家完成内部证据包之前。故事不得暗示暂停是由玩家单独促成。",
        GREEN,
    )
    action(doc, "把 E01 至 E06 按“直接记录、证人陈述、媒体转述、待专业复核”分类。")

    heading(doc, "场景四 环境资料对比")
    plain(doc, "资料库显示项目公开文件、中华白海豚保护资料和一张教学用地图。")
    production(
        doc,
        "地图、调查边界、目击点和坐标均为游戏化模拟数据。界面必须长期显示该标签。若制作组改用真实资料，须逐项补齐报告名称、版本、页码、坐标精度、原始来源和绘制方法。",
    )
    action(doc, "打开官方物种名录和中华白海豚保护行动计划。")
    colored(
        doc,
        "[研究结果]",
        "中华白海豚列入国家一级重点保护野生动物名录。相关保护资料把栖息地变化、人类活动和海上爆破等列为需要防范的威胁。一般风险知识可以说明为什么值得调查，但不能代替本案的损害证据。",
    )
    action(doc, "打开教学地图，依次叠加模拟调查范围、历史线索和玩家单次观察。")
    colored(
        doc,
        "[界面标签]",
        "游戏化模拟数据  仅用于证据判断教学  不构成真实物种分布图或环评结论",
        RED,
    )
    plain(doc, "部分模拟目击点落在简化调查范围外。系统没有自动弹出“环评违法”，而是新增三个问题。")
    colored(doc, "[待核实 Q01]", "公开环境评估文件实际调查了什么范围、采用了什么方法。")
    colored(doc, "[待核实 Q02]", "模拟目击点对应的原始资料是否存在，时间和坐标是否足够准确。")
    colored(doc, "[待核实 Q03]", "当前资料是否足以判断调查覆盖与白海豚活动之间存在实质缺口。")
    speaker(
        doc,
        "项目负责人",
        "居民没听说过一份评估，不代表评估不存在。文件是否公开存在、评估是否充分、公众是否实际知情，是三个不同问题。",
    )
    colored(doc, "[安全结论]", "当前对比只能提出复核问题，不能宣布评估无效或工程违法。", GREEN)

    heading(doc, "场景五 证据墙")
    plain(doc, "第二天，玩家站在证据墙前。左侧是主张，右侧是来源卡；每张卡都显示取得时间、来源类型和局限。")
    colored(
        doc,
        "[新任务]",
        "为每条主张选择真正相关的证据。允许保留“证据不足”，不要求所有问题都有答案。",
    )
    plain(doc, "主张一：公开资料显示，5月13日有关工程实施了两次炸礁作业。")
    action(doc, "链接 E01 公告与 E05 媒体对部门工作人员说法的转述。")
    colored(
        doc,
        "[审核]",
        "E01证明计划，E05支持“实施两次”的说法。结论必须保留来源说明，不能把E02施工船照片当成爆破证明。",
        GREEN,
    )
    plain(doc, "主张二：当天爆破造成中华白海豚受伤或死亡。")
    action(doc, "检查 E03 证词、E04 影像和 E06 伤亡信息。")
    choice(doc, "选择 A", "标记为“已经证实”。")
    colored(
        doc,
        "[系统结果]",
        "没有伤亡记录、专业检查或可支持因果关系的资料。过度断言会使报告可信度减2，并触发返工。",
        RED,
    )
    choice(doc, "选择 B", "标记为“证据不足，不能确认伤亡或因果关系”。")
    colored(doc, "[系统结果]", "报告保留风险，同时没有虚构结果。可信度加1。", GREEN)
    plain(doc, "主张三：围头湾出现中华白海豚，因此所有施工都必然违法。")
    colored(
        doc,
        "[审核]",
        "前半句需要可靠的物种和地点资料；后半句是法律结论，不能由单次目击推出。应改为“白海豚活动线索使工程风险与保护措施值得专业评估”。",
        RED,
    )
    plain(doc, "主张四：项目环境评估存在重大遗漏。")
    action(doc, "链接公开环境文件、教学地图和环保组织后续文章。")
    colored(
        doc,
        "[审核]",
        "环保组织提出了具体质疑，但那是当事一方的主张。教学地图不是现实证据；现有资料不能替代监管或法院判断。",
        GREEN,
    )
    plain(doc, "主张五：5月14日的暂停由玩家提交报告直接造成。")
    colored(
        doc,
        "[审核]",
        "暂停早于玩家的内部提交，且公开报道提到多方反映和监管介入。删除该因果关系。",
        RED,
    )
    action(doc, "查看最终证据清单，补齐每个文件的来源、时间、原始链接和局限说明。")

    heading(doc, "场景六 提交内部证据包")
    plain(doc, "时间卡：2025年5月17日 16:00")
    plain(doc, "内部报告表没有自动生成指控。四个固定栏目依次展开：已核实、来源所述、风险问题、仍未知。")
    plain(doc, "报告正文")
    paragraph(
        doc,
        [
            ("已核实  ", {"bold": True, "color": GREEN}),
            ("5月12日公开公告预告次日相关水下作业；玩家于5月13日傍晚从公共岸线记录到施工船，并保存了一段待物种复核的远距离影像。", {}),
        ],
    )
    paragraph(
        doc,
        [
            ("来源所述  ", {"bold": True, "color": PURPLE}),
            ("一名匿名渔民称当天下午听见两声闷响，但没有亲眼看见爆破。5月16日媒体报道转述部门工作人员说法，称5月13日进行了两次炸礁作业并于当晚暂停。", {}),
        ],
    )
    paragraph(
        doc,
        [
            ("风险问题  ", {"bold": True, "color": PURPLE}),
            ("中华白海豚是国家一级重点保护野生动物；水下爆破及栖息地改变属于需要专业评估的风险。本报告不据此断言个体已经受伤。", {}),
        ],
    )
    paragraph(
        doc,
        [
            ("仍未知  ", {"bold": True, "color": RED}),
            ("影像中的物种、动物数量与状态；声响的直接来源；项目评估和保护措施是否充分；后续作业安排及独立监测结果。", {}),
        ],
    )
    action(doc, "选择所有附件，运行“来源与隐私检查”，再点击提交给项目负责人。")
    colored(
        doc,
        "[系统]",
        "内部证据包已保存。它可能帮助组织决定是否继续核实或向有权限的部门反映。这份证据包不能证明违法，也不保证产生某个结果。",
    )

    heading(doc, "场景七 后续发展", page_break=True)
    plain(doc, "时间卡：2025年7月29日")
    plain(doc, "电脑收到一条公开信息更新。玩家打开自然之友发布的案件说明。")
    action(doc, "阅读原文，并把机构陈述与程序事实分开。")
    colored(
        doc,
        "[机构主张]",
        "自然之友称，项目在2025年6月中下旬又进行了新一轮炸礁，并对环境评估和保护措施提出质疑。这些内容代表该环保组织的调查与法律主张，不是法院已经确认的事实。",
    )
    colored(
        doc,
        "[程序事实]",
        "自然之友称，泉州市中级人民法院立案庭于2025年7月28日签收其提交的生态破坏民事公益诉讼材料。签收材料不等于正式立案，更不等于胜诉。",
        GREEN,
    )
    choice(doc, "选择 A", "把新闻卡写成“法院已经认定工程破坏白海豚栖息地”。")
    colored(
        doc,
        "[系统结果]",
        "材料签收阶段没有实体裁判。该标题错误，调查可信度减2。",
        RED,
    )
    choice(doc, "选择 B", "写成“环保组织提交公益诉讼材料，法院是否立案仍需后续信息”。")
    colored(doc, "[系统结果]", "程序状态和实体争议已分开。可信度加1。", GREEN)
    plain(doc, "时间卡：2026年8月26日")
    colored(
        doc,
        "[后续报道]",
        "一篇媒体后续报道写道，法院于2025年8月以社会组织主体资格为由未予立案。这是程序结果，不是对工程是否合法、评估是否充分或生态损害是否发生的实体判断。",
    )
    production(
        doc,
        "该诉讼状态来自2026年8月的媒体后续报道。发行前需再次核对法院或当事机构是否公布新信息；如有变化，只更新数据层，不改写玩家已完成的2025年调查。",
    )

    heading(doc, "场景八 回到海边")
    plain(doc, "时间卡：2026年9月")
    plain(doc, "玩家再次站在同一处公共岸线。码头设施已经成形，一艘船从远处经过。")
    colored(
        doc,
        "[现实更新]",
        "泉州市海洋与渔业局2026年3月的公开信息将11号、12号泊位工程称为在建海洋工程；南安市人力资源和社会保障局2026年5月的工资专户销户公示则称该工程已竣工。两项信息对应不同时间和行政事项，均不能说明生态争议已经得到实体裁判。",
    )
    action(doc, "架好相机，填写新的观察日期、位置、天气和能见度。")
    plain(doc, "水下传来清晰的回声定位声。画面切到浑浊的蓝绿色海水，一道浅色背影从远处掠过。")
    production(
        doc,
        "本镜头是象征性结尾，不代表真实个体、真实日期或监测结论。不得加入搁浅、流血、死亡或施工直接击中白海豚的画面。",
    )
    plain(doc, "字幕：有些声音，人类未必听得见。证据要做的，是让风险被准确看见。")
    colored(
        doc,
        "[章节结论]",
        "保护海洋动物不只包括救助受伤个体，也包括长期了解并保护它们依赖的栖息地。一次调查不能给出所有答案，但准确记录比夸大结论更有用。",
        GREEN,
    )
    colored(doc, "[奖励]", "获得真相纹章碎片。")
    action(doc, "点击“结束本章”。")

    heading(doc, "结尾信息卡")
    paragraph(
        doc,
        [
            ("中华白海豚  ", {"bold": True, "color": GREEN}),
            ("Sousa chinensis", {"italic": True, "color": GREEN}),
        ],
    )
    colored(doc, "[事实]", "中华白海豚是国家一级重点保护野生动物。", GREEN)
    colored(
        doc,
        "[方法]",
        "记录时间、地点、原始来源和局限；把亲眼观察与推测分开；不为取得证据进入危险或非公开区域。",
        GREEN,
    )
    colored(
        doc,
        "[边界]",
        "施工公告证明计划，现场照片证明画面中的对象，证人陈述证明对方说过什么。动物伤亡、生态影响和违法性需要相应的专业证据与有权判断。",
        GREEN,
    )
    plain(doc, "感谢游玩")

    heading(doc, "制作组来源说明", page_break=True)
    production(
        doc,
        "以下来源用于校准现实事件、物种信息和截至2026年9月9日的后续状态。游戏正文已按来源类型标明“官方文件”“媒体转述”“机构主张”或“后续报道”。",
    )
    source(
        doc,
        "农业部 中华白海豚保护行动计划 2017至2026年",
        "https://yyj.moa.gov.cn/bjwj/201904/t20190419_6197438.htm",
        "用于物种保护背景和海上工程风险的一般说明，不用于证明本案已经发生具体损害。",
    )
    source(
        doc,
        "农业农村部 国家重点保护野生动物名录",
        "https://yyj.moa.gov.cn/gzdt/202102/P020210205564727206658.pdf",
        "用于确认中华白海豚的国家一级保护等级。",
    )
    source(
        doc,
        "福建省交通运输厅 泉州港围头湾港区石井作业区11号12号泊位工程初步设计批复",
        "https://jtyst.fujian.gov.cn/qzg/zfxxgk/zfxxgkml/xzxkyqr/202301/t20230112_6093650.html",
        "用于工程名称、位置与建设规模的背景核对。",
    )
    source(
        doc,
        "福建省交通运输厅 关于同意有关泊位工程爆破活动的批复",
        "https://jtyst.fujian.gov.cn/qzg/zfxxgk/zfxxgkml/xzxkyqr/202505/t20250520_6916289.html",
        "官方行政许可页面。只说明许可文件内容，不单独证明某次具体作业已经实施或造成生态影响。",
    )
    source(
        doc,
        "澎湃新闻 泉州围头湾一港区项目炸礁被指影响中华白海豚 官方称已叫停重新评估",
        "https://m.thepaper.cn/newsDetail_forward_30827546",
        "2025年5月16日报道。用于明确“两次作业、当晚暂停、次日到场评估”属于媒体对部门工作人员说法的转述。",
    )
    source(
        doc,
        "红星新闻 围头湾炸礁施工被指或伤害中华白海豚",
        "https://www.sohu.com/a/895517942_116237",
        "2025年5月15日报道的转载页面。用于“当时暂未收到伤亡报告”的有限信息，不能据此推出不存在影响。",
    )
    source(
        doc,
        "自然之友 起诉 用公益诉讼拯救中华白海豚",
        "https://www.fon.org.cn/action/way/content/558",
        "2025年7月29日发布。用于该组织关于6月作业、环境评估问题与诉讼请求的主张，以及7月28日立案庭签收材料这一程序信息。不得写成法院认定。",
    )
    source(
        doc,
        "福建省水产研究所 泉州围头湾中华白海豚及其栖息地保护研讨会",
        "https://www.fjscs.ac.cn/notice_details-714.html",
        "用于确认围头湾白海豚分布、繁殖成长、栖息地变化、周边建设及影响评价要求属于持续研究与讨论的问题。",
    )
    source(
        doc,
        "泉州市海洋与渔业局 海洋工程联合检查",
        "https://hyj.quanzhou.gov.cn/zwgk/zfxxgkml/gzdt/202603/t20260309_3272571.htm",
        "2026年3月7日公开信息。用于确认2026年3月官方仍将该泊位工程列为在建海洋工程，并开展检查和环保指导。",
    )
    source(
        doc,
        "南安市人力资源和社会保障局 泊位项目工资专户销户公示",
        "https://www.nanan.gov.cn/zwgk/xwzx/tzgg/202605/t20260522_3293824.htm",
        "2026年5月22日公开信息。公示称该工程已竣工并提交工资支付资料；该表述服务于工资专户销户，不等于生态争议已有裁判结论。",
    )
    source(
        doc,
        "对话地球 社会组织为何无法以公益诉讼保护白海豚",
        "https://dialogue.earth/zh/9/60148831/",
        "2026年8月26日后续报道。用于说明媒体所报的未予立案及主体资格问题；这不是项目争议的实体裁判。",
    )

    heading(doc, "发行前检查清单")
    plain(doc, "1. 重新核对公告、工程状态、诉讼程序和对外联系渠道，并更新信息日期。")
    plain(doc, "2. 教学地图若保留模拟数据，标签必须持续可见；若替换真实数据，补齐报告版本、页码、坐标精度和授权。")
    plain(doc, "3. 不把施工船照片、声响证词、单次目击或机构主张升级为动物伤亡、因果关系或违法结论。")
    plain(doc, "4. 方言、村落建筑和渔业劳动细节由当地顾问审核；不得用错别字模拟口音。")
    plain(doc, "5. Unity中的事实卡、来源链接和时间线使用可更新配置，不写死在场景脚本中。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        os.chmod(OUTPUT, stat.S_IREAD | stat.S_IWRITE)
    doc.save(OUTPUT)

    check = Document(OUTPUT)
    full_text = "\n".join(p.text for p in check.paragraphs)
    banned = [
        "mississippi",
        "观察战",
        "清晰人生",
        "手上个体",
        "晚12:30",
        "上个拜三",
        "他们的存亡",
        "官方通报，可证明爆破行为",
        "文件件",
    ]
    for token in banned:
        assert token.lower() not in full_text.lower(), f"Unsafe phrase remains: {token}"
    required = [
        "2025年5月12日",
        "2025年5月13日",
        "2025年5月16日",
        "2025年5月17日",
        "2025年7月29日",
        "2026年8月26日",
        "2026年9月",
        "游戏化模拟数据",
        "媒体对部门工作人员说法的转述",
        "不等于正式立案",
    ]
    for token in required:
        assert token in full_text, f"Required content missing: {token}"
    timeline_tokens = [
        "时间卡：2025年5月12日 19:30",
        "时间卡：2025年5月13日 18:35",
        "时间卡：2025年5月16日 20:10",
        "时间卡：2025年5月17日 16:00",
        "时间卡：2025年7月29日",
        "时间卡：2026年8月26日",
        "时间卡：2026年9月",
    ]
    positions = [full_text.index(token) for token in timeline_tokens]
    assert positions == sorted(positions), "Narrative dates are not in chronological order"
    assert len(check.paragraphs) >= 100
    with zipfile.ZipFile(OUTPUT) as archive:
        xml = archive.read("word/document.xml")
        assert b"<w:ins" not in xml and b"<w:del" not in xml
        assert "word/comments.xml" not in archive.namelist()
    print(OUTPUT)
    print(f"Paragraphs: {len(check.paragraphs)}")


if __name__ == "__main__":
    build()
