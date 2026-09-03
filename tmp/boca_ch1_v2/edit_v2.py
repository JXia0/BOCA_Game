from copy import deepcopy
from pathlib import Path
import os
import shutil
import stat

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\GameProjects\BOCA_Game\output\docx\BOCA 游戏剧本_玳瑁古筝义甲_修改版.docx")
OUTPUT = Path(r"C:\GameProjects\BOCA_Game\output\docx\BOCA 游戏剧本_玳瑁古筝义甲_修改版_v2.docx")

PURPLE = RGBColor(0xA0, 0x2B, 0x93)
RED = RGBColor(0xEE, 0x00, 0x00)
GREEN = RGBColor(0x00, 0xB0, 0x50)
GRAY = RGBColor(0x66, 0x66, 0x66)


def find(doc, exact):
    matches = [p for p in doc.paragraphs if p.text == exact]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph for {exact!r}, found {len(matches)}")
    return matches[0]


def clear_runs(paragraph):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_parts(paragraph, parts):
    clear_runs(paragraph)
    for text, fmt in parts:
        run = paragraph.add_run(text)
        run.bold = fmt.get("bold", False)
        run.italic = fmt.get("italic", False)
        if "color" in fmt:
            run.font.color.rgb = fmt["color"]
        if fmt.get("highlight") == "yellow":
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif fmt.get("highlight") == "turquoise":
            run.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
        if "size" in fmt:
            run.font.size = Pt(fmt["size"])
    return paragraph


def set_plain(paragraph, text):
    return set_parts(paragraph, [(text, {})])


def insert_after(paragraph, parts, *, space_after=8):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    result = Paragraph(new_p, paragraph._parent)
    result.paragraph_format.space_after = Pt(space_after)
    set_parts(result, parts)
    return result


def delete_paragraph(paragraph):
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        os.chmod(OUTPUT, stat.S_IREAD | stat.S_IWRITE)
    shutil.copy2(SOURCE, OUTPUT)
    os.chmod(OUTPUT, stat.S_IREAD | stat.S_IWRITE)
    doc = Document(OUTPUT)

    doc.core_properties.title = "BOCA - Hawksbill Guzheng Finger Pick Case (Revised Script v2)"

    set_parts(find(doc, "Revised gameplay script based on the original student draft"), [
        ("Revised gameplay script based on the original student draft - v2", {"italic": True, "color": GRAY, "size": 10})
    ])

    # Remove the early warning. The player and friend should not begin the story
    # already suspecting a wildlife offence.
    delete_paragraph(find(doc, "[SYSTEM] A seller's description is a claim, not proof of the material or its legal status."))

    set_plain(find(doc, "Two routes appear. Neither route requires the player to purchase the suspicious product."),
              "Two routes appear. At this point, the player is shopping normally and does not yet know what “玳瑁” means.")
    set_parts(find(doc, "Choice A: Open the advertised “玳瑁古筝义甲” listing to inspect its description."), [
        ("Choice A: ", {"bold": True, "highlight": "turquoise"}),
        ("Buy the seller-recommended “玳瑁古筝义甲” set.", {}),
    ])
    set_parts(find(doc, "Choice B: Purchase a clearly labelled synthetic set and leave the suspicious listing unopened."), [
        ("Choice B: ", {"bold": True, "highlight": "turquoise"}),
        ("Buy a clearly labelled synthetic set instead.", {}),
    ])

    set_plain(find(doc, "The suspicious product page fills the phone screen. The listing repeats “genuine hawksbill” but shows no visible approval document, traceability information, or special identification mark."),
              "The product page presents “genuine hawksbill” as a premium traditional material recommended for serious students. The player understands it only as the name of a material and does not examine wildlife documents or traceability information.")
    set_parts(find(doc, "Open the seller chat. Ask what the material is and whether the seller can provide lawful approval and traceability information."), [
        ("Open the seller chat. Ask whether the set is comfortable, durable, and suitable for the upcoming examination.", {"highlight": "yellow"})
    ])
    set_parts(find(doc, "SELLER: It is real hawksbill. Old material. Everyone uses it. You do not need to worry about paperwork."), [
        ("SELLER: ", {"bold": True}),
        ("It is genuine hawksbill, a traditional material. Many teachers and advanced players prefer it. This size should suit you.", {}),
    ])
    set_parts(find(doc, "[SYSTEM] The seller has repeated the claim but has not provided evidence. Do not place an order."), [
        ("[SYSTEM] ", {"bold": True, "color": PURPLE}),
        ("The player has not researched the word “玳瑁” and does not interpret the purchase as risky or unusual.", {"color": PURPLE}),
    ])
    purchase_action = find(doc, "Close the payment window without purchasing.")
    set_parts(purchase_action, [("Confirm the purchase.", {"highlight": "yellow"})])
    anchor = insert_after(purchase_action, [("Several days later, the package arrives. The material name “玳瑁” is printed on the label inside the box.", {})])
    anchor = insert_after(anchor, [("Open the package and try on the finger picks.", {"highlight": "yellow"})], space_after=3)
    anchor = insert_after(anchor, [("The picks feel firm and unfamiliar. Out of simple curiosity about what the material name means, the player decides to look it up online.", {})])

    set_plain(find(doc, "The player selects a synthetic set with a clear material description. Later that evening, a friend sends a screenshot of the suspicious “玳瑁” listing and asks whether the product is really made from a protected animal."),
              "The player selects a synthetic set with a clear material description. Later that evening, a friend casually sends a screenshot of a “玳瑁” set and writes: “My teacher recommended this material. People say it feels good to play with - have you tried it?”")
    set_parts(find(doc, "Save the friend's screenshot and open the suspicious listing without purchasing it."), [
        ("Save the friend's screenshot so the player can remember the unfamiliar material name.", {"highlight": "yellow"})
    ])
    set_parts(find(doc, "[SYSTEM] Both routes now continue to the same research scene."), [
        ("[SYSTEM] ", {"bold": True, "color": PURPLE}),
        ("In both routes, curiosity - not suspicion - leads the player to search for the meaning of “玳瑁”.", {"color": PURPLE}),
    ])

    set_plain(find(doc, "The computer screen lights the room. The player places the marketplace window beside an official wildlife-information page."),
              "The computer screen lights the room. In Route A, the product label rests beside the keyboard. In Route B, the friend's screenshot remains open on the phone.")
    set_parts(find(doc, "Drag the keywords “玳瑁”, “国家重点保护野生动物名录”, and “野生动物制品 专用标识” into the search bar."), [
        ("Drag the unfamiliar word “玳瑁” into the search bar.", {"highlight": "yellow"})
    ])
    result_one = find(doc, "RESEARCH RESULT: The hawksbill turtle (玳瑁, Eretmochelys imbricata) is listed as a Class I national key protected wild animal in China.")
    set_parts(result_one, [
        ("SEARCH RESULT: ", {"bold": True, "color": PURPLE}),
        ("“玳瑁” is the Chinese name for the hawksbill turtle (Eretmochelys imbricata), a marine animal whose shell has historically been used to make decorative objects and small accessories.", {"color": PURPLE}),
    ])
    insert_after(result_one, [("Open official sources and search “国家重点保护野生动物名录” and “野生动物制品 专用标识”.", {"highlight": "yellow"})], space_after=3)
    set_parts(find(doc, "RESEARCH RESULT: National rules generally prohibit the sale, purchase, and use of products made from nationally protected wildlife. Special lawful uses require approval and traceability requirements under the applicable rules."), [
        ("RESEARCH RESULT: ", {"bold": True, "color": PURPLE}),
        ("The hawksbill turtle is listed as a Class I national key protected wild animal in China. National rules generally prohibit the sale, purchase, and use of products made from nationally protected wildlife; special lawful uses require approval and traceability under the applicable rules.", {"color": PURPLE}),
    ])
    set_parts(find(doc, "IMPORTANT DISTINCTION: The listing is suspicious, but the player cannot prove from an online photograph alone that the object is genuine hawksbill. The report must separate the seller's claim from verified facts."), [
        ("IMPORTANT DISTINCTION: ", {"bold": True, "color": RED}),
        ("The discovery makes the listing a potential wildlife-trade lead, but an online photograph and seller claim still cannot prove that the item is genuine hawksbill. The report must separate claims from verified facts.", {"color": RED}),
    ])

    set_parts(find(doc, "[SYSTEM] New task: preserve the listing and conversation without making another transaction or confronting the seller."), [
        ("[SYSTEM] ", {"bold": True, "color": PURPLE}),
        ("New task: preserve what already happened. Route A includes a completed purchase; Route B includes only a recommendation and listing. Do not make another transaction or confront the seller.", {"color": PURPLE}),
    ])
    set_parts(find(doc, "Capture the seller conversation showing the material claim and the request for approval or traceability information."), [
        ("Capture the seller conversation showing how the material was described and recommended.", {"highlight": "yellow"})
    ])

    set_parts(find(doc, "Optional route: The player selects “I already bought a similar item before I understood the issue.”"), [
        ("ROUTE A - PURCHASE RECORDS", {"bold": True, "color": PURPLE}),
    ])
    set_parts(find(doc, "The interface adds the order number, payment record, logistics information, packaging, and labels to the evidence checklist."), [
        ("Add the order number, payment record, logistics information, full packaging, and labels to the evidence folder.", {"highlight": "yellow"})
    ])
    set_parts(find(doc, "[PRODUCTION NOTE] The optional route represents a purchase that occurred before the player understood the concern. The game must never require or reward a new purchase of a suspected wildlife product."), [
        ("[PRODUCTION NOTE] ", {"bold": True, "color": GRAY}),
        ("This route reflects the real 2023 case pattern: the student bought the recommended set, received it, then searched the unfamiliar material name out of curiosity and reported the discovery. The game must show that the player had no informed suspicion before buying.", {"italic": True, "color": GRAY}),
    ])
    route_a_note = find(doc, "[PRODUCTION NOTE] This route reflects the real 2023 case pattern: the student bought the recommended set, received it, then searched the unfamiliar material name out of curiosity and reported the discovery. The game must show that the player had no informed suspicion before buying.")
    anchor = insert_after(route_a_note, [("ROUTE B - NO PURCHASE", {"bold": True, "color": PURPLE})], space_after=3)
    anchor.paragraph_format.keep_with_next = True
    anchor.paragraph_format.keep_together = True
    insert_after(anchor, [("Attach the friend's screenshot and the listing records. Do not buy the product to obtain additional evidence.", {"highlight": "yellow"})])

    set_parts(find(doc, "If the optional prior-purchase route was selected, attach the order, payment, logistics, packaging, and label records."), [
        ("For Route A, attach the order, payment, logistics, packaging, and label records. For Route B, attach the friend's screenshot and the listing records.", {"highlight": "yellow"})
    ])
    set_plain(find(doc, "DRAFT REPORT"), "DRAFT REPORT - ROUTE A (PURCHASED ITEM)")
    report_a = find(doc, "At 19:42 on 18 June 2025, I found an online listing advertising guzheng finger picks as “genuine hawksbill.” The seller repeated that description in the attached conversation but did not provide approval or traceability information when asked. I have not independently verified the material. The attached files preserve the listing, seller statements, page address, store account, and relevant official species information. I am submitting these records for assessment.")
    set_plain(report_a, "At 19:42 on 18 June 2025, I purchased guzheng finger picks advertised and recommended as “genuine hawksbill.” After receiving the item, I searched the material name out of curiosity and learned that the hawksbill turtle is nationally protected. I have stopped using the item. I have not independently verified its material. The attached files preserve the listing, seller statements, order, payment, logistics, packaging, labels, and relevant official species information. I am submitting these records for assessment and instructions on handling the item.")
    anchor = insert_after(report_a, [("DRAFT REPORT - ROUTE B (LISTING ONLY)", {"bold": True})], space_after=3)
    insert_after(anchor, [("At 20:06 on 18 June 2025, a friend sent me an online listing for guzheng finger picks advertised as “genuine hawksbill.” My friend presented it as a normal recommendation and did not express suspicion. After searching the unfamiliar material name, I learned that the hawksbill turtle is nationally protected. I did not purchase the product and have not independently verified its material. The attached files preserve the screenshot, listing, page address, store account, and relevant official species information. I am submitting these records for assessment.", {})])

    set_plain(find(doc, "The synthetic finger picks rest beside the strings. The player puts them on and begins to play. The sound is ordinary, clear, and enough."),
              "A safe synthetic set rests beside the strings. The player puts it on and begins to play. The sound is ordinary, clear, and enough.")
    set_parts(find(doc, "HAWKSBILL TURTLE / 玳瑁\nThe hawksbill turtle is a protected marine animal. A product description or visual resemblance cannot by itself establish what a product is made from. Avoid suspicious purchases, preserve accurate evidence, distinguish claims from verified facts, and use current official reporting channels."), [
        ("HAWKSBILL TURTLE / 玳瑁", {"bold": True, "color": GREEN}),
        ("\nThe hawksbill turtle is a protected marine animal. A product description or visual resemblance cannot by itself establish what a product is made from. When an unfamiliar animal-derived material appears in a product description, pause and check official sources. If a concern is discovered after purchase, stop using the item, preserve accurate records, and seek instructions through verified reporting channels.", {"color": GREEN}),
    ])

    source_heading = find(doc, "Source Notes for the Production Team")
    source_heading.paragraph_format.page_break_before = True
    anchor = insert_after(source_heading, [
        ("REAL CASE INSPIRATION: ", {"bold": True}),
        ("In March 2023, a guzheng student in Xuzhou bought a seller-recommended hawksbill finger-pick set for an examination. After delivery, the student searched the material out of curiosity, discovered the wildlife-protection concern, and reported it to police.", {}),
    ], space_after=3)
    insert_after(anchor, [("https://www.jsjc.gov.cn/zl/6736/202506/t20250624_189512.shtml", {"color": PURPLE, "size": 9})])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
