from pathlib import Path
import sys
from copy import deepcopy
from docx import Document
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'boca_ch1_v2'))
from edit_v2 import find, set_plain, set_parts, insert_after, PURPLE, GRAY

ROOT = Path(r'C:\GameProjects\BOCA_Game')
SOURCE = ROOT / 'output/docx/BOCA 游戏剧本_玳瑁古筝义甲_修改版_v2.docx'
OUTPUT = ROOT / 'output/docx/BOCA 游戏剧本_玳瑁古筝义甲_修改版_v3.docx'
doc = Document(SOURCE)

def starts(prefix):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(matches) == 1, (prefix, len(matches))
    return matches[0]

def comment(p, text):
    doc.add_comment(p.runs, text=text, author='BOCA Revision', initials='BR')

set_parts(starts('Revised gameplay script'), [
    ('Revised gameplay script based on the original student draft - v3', {'italic': True, 'color': GRAY, 'size': 10})
])
doc.core_properties.title = 'BOCA - Hawksbill Guzheng Finger Pick Case (Revised Script v3)'

set_parts(starts('SELLER:'), [
    ('SELLER: ', {'bold': True}),
    ('It is genuine hawksbill, a traditional material. Customers often choose it for its firm feel. This size should suit you.', {})
])
friend = starts('The player selects a synthetic set')
set_plain(friend, 'The player selects a synthetic set with a clear material description. Later that evening, a friend casually sends a photo of the new finger picks they have just bought, followed by the listing screenshot. The message reads: “I just got this set! Tried it during practice today - it feels pretty good. The listing calls it ‘玳瑁’. Have you tried this kind?”')
comment(friend, 'Changed at your request: the friend has bought and tried the set and shares a personal impression. No teacher recommendation and no initial suspicion. The friend does not know the material is a protected animal product.')

set_parts(starts('[SYSTEM] New task:'), [
    ('[SYSTEM] ', {'bold': True, 'color': PURPLE}),
    ('New task: preserve what already happened. In Route A, the player has bought the item. In Route B, the friend has bought a set, but the player has only the shared photo, messages, and listing. Do not make another transaction or confront the seller.', {'color': PURPLE})
])
set_parts(starts('Capture the seller conversation'), [
    ('Route A: capture the existing seller conversation. Route B: save the friend\'s original messages, photo, and listing screenshot; the player has no seller chat of their own.', {'highlight': 'yellow'})
])
set_plain(starts('Seller chat -'), "Conversation records - Route A: the seller's material claim; Route B: the friend's purchase and use account. Neither independently verifies the material.")
set_parts(find(doc, 'ROUTE B - NO PURCHASE'), [('ROUTE B - PLAYER HAS NOT PURCHASED', {'bold': True, 'color': PURPLE})])
set_parts(starts("Attach the friend's screenshot and the listing records."), [
    ("Attach the friend's shared messages, photo, screenshot, and the listing records. Do not buy the product to obtain additional evidence.", {'highlight': 'yellow'})
])
set_parts(starts('Attach the listing screenshots, seller-chat screenshots'), [
    ('Attach the listing screenshots, page address, timestamps, official wildlife source, and the conversation records available in the selected route.', {'highlight': 'yellow'})
])
set_parts(starts('For Route A, attach the order'), [
    ("For Route A, attach the player's order, payment, logistics, packaging, and label records. For Route B, attach only the messages and images the friend has shared; do not invent or automatically access the friend's private order records.", {'highlight': 'yellow'})
])
set_plain(starts('At 20:06 on 18 June 2025'), 'At 20:06 on 18 June 2025, a friend sent me a photo and an online listing for finger picks advertised as “genuine hawksbill.” My friend said they had bought the set and found it good to use. After searching the unfamiliar material name, I learned that the hawksbill turtle is nationally protected. I did not purchase the product and have not independently verified its material or my friend\'s purchase. The attached files preserve the shared messages, photo, listing screenshot, page address, store account, and relevant official species information. I am submitting these records for assessment.')

# Insert the requested seller cutaway between the personal follow-up and the information card.
# Its fictional, audience-only viewpoint must not become evidence accessible to the protagonist.
ending = find(doc, 'Ending Information Card')
paragraphs = doc.paragraphs
ending_index = next(i for i,p in enumerate(paragraphs) if p._p is ending._p)
anchor = paragraphs[ending_index - 1]
heading_template = starts('Scene 6:')

def add(text, fmt=None, heading=False, page_break=False):
    global anchor
    anchor = insert_after(anchor, [(text, fmt or {})])
    if text.startswith(('BUYER: ', 'SELLER: ')):
        speaker, dialogue = text.split(': ', 1)
        set_parts(anchor, [(speaker + ': ', {'bold': True}), (dialogue, {})])
    if heading:
        anchor.style = heading_template.style
        if heading_template._p.pPr is not None:
            anchor._p.remove(anchor._p.pPr)
            anchor._p.insert(0, deepcopy(heading_template._p.pPr))
        anchor.paragraph_format.keep_with_next = True
    anchor.paragraph_format.page_break_before = page_break
    return anchor

new_heading = add('Scene 7: The Seller - Elsewhere, Around the Same Time', {'bold': True}, heading=True)
comment(new_heading, 'Added at your request, following the second screenshot: seller at a desktop, incoming orders and payments, then a later police intervention and handcuffs. Kept the existing fictional 2025 chronology. No 300,000 profit counter or September 2023 date is asserted as fact.')
add('[PRODUCTION NOTE] This is an audience-only, fictionalized cutaway inspired by the case, not footage viewed by the protagonist or evidence in their report. Use a fictional shop interface. The scene is not an exact reconstruction of the real arrest.', {'italic': True, 'color': GRAY})
add('The last guzheng note fades into the clicking of a keyboard. The camera cuts to a cramped packing room, behind a man at his desk. The shop avatar matches the listing seen earlier. His desktop shows a buyer conversation; labelled packets of finger picks lie beside the keyboard.')
add('He types to another anonymous buyer. On the right, new orders join a queue. Payment-received notices alternate with outgoing stock and shipping entries. Another payment notification sounds a few seconds later. No total profit figure is shown.')
add('BUYER: Is this the same material as the one in your listing?')
add('SELLER: Yes, genuine hawksbill. It is in stock. I can send it tomorrow.')
add('He sends the reply, seals another packet, and places it on a growing stack. A new order arrives before he finishes. The camera stays on the hands, packages, and screen; there is no villain monologue.')
add('Press Next.', {'highlight': 'yellow'})

add('Scene 8: Police Intervention - Months Later', {'bold': True}, heading=True)
add('The screen fades to black. A time card reads: “MONTHS LATER”. Brief, silent shots show investigators comparing transaction records, a specialist examining seized material, and an evidence folder being closed. These shots establish a separate investigation over time.')
add('Cut back to the packing room. Police officers are present. The seller has stopped typing; an unfinished buyer reply remains on the monitor. In close-up, an officer closes handcuffs around his wrists. The metal click cuts through the silence.')
add('The camera holds on the abandoned keyboard. A payment notification sounds once more, but nobody answers. Fade to black.')
add('[PRODUCTION NOTE] The handcuff shot depicts the fictional seller being taken into custody, not a conviction or the automatic result of one report. Both routes reach this epilogue; the protagonist never controls the seller or the police.', {'italic': True, 'color': GRAY})
add('Press Next to view the ending information card.', {'highlight': 'yellow'})
ending.paragraph_format.keep_with_next = True

source_note = starts('REAL CASE INSPIRATION:')
set_parts(source_note, [
    ('REAL CASE INSPIRATION: ', {'bold': True}),
    ('In March 2023, a guzheng student in Xuzhou bought a seller-recommended hawksbill finger-pick set for an examination. After delivery, the student searched the material out of curiosity, discovered the wildlife-protection concern, and reported it to police. The cited report also describes later investigation and court proceedings. This script keeps a fictional 2025 setting; the friend, dialogue, order notifications, investigation montage, and handcuff staging are dramatized, not verified details of the real incident.', {})
])

doc.save(OUTPUT)
check = Document(OUTPUT)
body = '\n'.join(p.text for p in check.paragraphs)
assert 'My teacher recommended' not in body
assert 'Many teachers' not in body
assert 'MONTHS LATER' in body and 'handcuffs' in body
assert len(check.comments) == 2
print(OUTPUT)
print(f'Paragraphs: {len(check.paragraphs)}; anchored revision comments: {len(check.comments)}')
