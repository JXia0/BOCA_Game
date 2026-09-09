import json
import sys
import zipfile
from pathlib import Path
from lxml import etree

from docx import Document
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")


def run_props(run):
    rpr = run._element.rPr
    highlight = None
    color = None
    if rpr is not None:
        hl = rpr.find(qn("w:highlight"))
        if hl is not None:
            highlight = hl.get(qn("w:val"))
        c = rpr.find(qn("w:color"))
        if c is not None:
            color = c.get(qn("w:val"))
    return {"text": run.text, "highlight": highlight, "color": color, "bold": run.bold}


for raw in sys.argv[1:]:
    path = Path(raw)
    doc = Document(path)
    data = {
        "path": str(path),
        "sections": [
            {
                "width": s.page_width,
                "height": s.page_height,
                "top": s.top_margin,
                "bottom": s.bottom_margin,
                "left": s.left_margin,
                "right": s.right_margin,
            }
            for s in doc.sections
        ],
        "paragraphs": [],
        "tables": [],
        "zip_flags": {},
    }
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            data["paragraphs"].append(
                {"i": i, "style": p.style.name if p.style else None, "text": p.text, "runs": [run_props(r) for r in p.runs if r.text]}
            )
    for ti, table in enumerate(doc.tables):
        data["tables"].append({"i": ti, "rows": [[cell.text for cell in row.cells] for row in table.rows]})
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        data["zip_flags"] = {
            "comments": "word/comments.xml" in names,
            "footnotes": "word/footnotes.xml" in names,
            "endnotes": "word/endnotes.xml" in names,
            "tracked_insertions": b"<w:ins" in zf.read("word/document.xml"),
            "tracked_deletions": b"<w:del" in zf.read("word/document.xml"),
        }
        root = etree.fromstring(zf.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        data["xml_paragraphs"] = []
        for i, p in enumerate(root.xpath("//w:body//w:p", namespaces=ns)):
            texts = p.xpath(".//w:t/text()", namespaces=ns)
            deleted = p.xpath(".//w:delText/text()", namespaces=ns)
            if texts or deleted:
                data["xml_paragraphs"].append({"i": i, "text": "".join(texts), "deleted": "".join(deleted)})
    print(json.dumps(data, ensure_ascii=False, indent=2))
